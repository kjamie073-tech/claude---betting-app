"""Bet tracker spreadsheet (Excel) with live ROI formulas.

The workbook has a Settings sheet (bank, unit size, staking rules), a Bets
sheet you add a row to for every bet, and a Summary sheet that works out
profit, ROI, strike rate, closing-line value and results by month, bet type
and market. Everything is formulas, so it updates as you fill it in, in Excel,
Google Sheets or Numbers.

Python helpers here let Claude add bets and settle them for you, and read the
sheet back to report ROI.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

MAX_ROWS = 2000
BET_TYPES = ["Single", "Player", "Builder"]
MARKETS = ["Result", "Goals", "BTTS", "Corners", "Cards", "Shots", "Player", "Builder", "Other"]
RESULTS = ["Pending", "Won", "Lost", "Void", "Cash out"]

# Bets sheet columns: (header, width)
COLUMNS = [
    ("Date", 12), ("Match", 26), ("Bookmaker", 12), ("Type", 9), ("Market", 10),
    ("Selection / legs", 44), ("Odds", 7), ("Model chance", 9), ("Fair odds", 8),
    ("Edge", 8), ("Suggested stake", 10), ("Stake", 8), ("Result", 9),
    ("Cash-out return", 10), ("Return", 9), ("Profit", 9), ("Closing odds", 9),
    ("CLV", 8), ("Expected profit", 10), ("Bank after", 11), ("Notes", 30),
]
COL = {name: get_column_letter(i + 1) for i, (name, _) in enumerate(COLUMNS)}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF7D6")
THIN = Side(style="thin", color="D0D0D0")


def _rng(col: str) -> str:
    return f"Bets!${COL[col]}$2:${COL[col]}${MAX_ROWS + 1}"


def create(path: str | Path, starting_bank: float = 500.0, start: dt.date | None = None) -> Path:
    path = Path(path)
    start = start or dt.date.today()
    wb = Workbook()

    # ------------------------------------------------------------ Settings
    s = wb.active
    s.title = "Settings"
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 14
    s.column_dimensions["C"].width = 70
    s["A1"] = "Staking settings"
    s["A1"].font = Font(bold=True, size=14)
    settings = [
        ("Starting bank (£)", starting_bank, "Money set aside purely for betting, that you can afford to lose.", "StartBank"),
        ("Start date", start, "", "StartDate"),
        ("Unit size (% of bank)", 0.01, "1 unit = 1% of the bank. Re-set monthly to 1% of the current bank.", "UnitPct"),
        ("Kelly fraction: singles", 0.25, "Quarter Kelly on result, goals, corners and cards singles.", "KellySingle"),
        ("Kelly fraction: player bets and builders", 0.125, "Eighth Kelly: these probabilities are less certain.", "KellyBuilder"),
        ("Minimum edge: singles", 0.03, "Result, goals, corners, cards singles. Skip anything with a smaller edge.", "MinEdgeSingle"),
        ("Minimum edge: player singles", 0.05, "Scorer, shots, cards for one player.", "MinEdgePlayer"),
        ("Minimum edge: builders", 0.08, "Builders need a bigger cushion for model error.", "MinEdgeBuilder"),
        ("Max stake per single (units)", 2, "", "CapSingle"),
        ("Max stake per builder (units)", 1, "", "CapBuilder"),
        ("Max stakes per match day (units)", 5, "Total across all bets that day.", "CapDay"),
        ("Stop and review if bank falls by", 0.4, "Pause and check the model's record before betting on.", "StopLoss"),
    ]
    for i, (label, val, note, name) in enumerate(settings, start=3):
        s[f"A{i}"] = label
        s[f"B{i}"] = val
        s[f"C{i}"] = note
        s[f"B{i}"].fill = INPUT_FILL
        if isinstance(val, float) and val < 1:
            s[f"B{i}"].number_format = "0.0%"
        if isinstance(val, dt.date):
            s[f"B{i}"].number_format = "dd mmm yyyy"
        if name == "StartBank":
            s[f"B{i}"].number_format = "£#,##0.00"
        wb.defined_names[name] = DefinedName(name, attr_text=f"Settings!$B${i}")
    r = 3 + len(settings) + 1
    s[f"A{r}"] = "Current bank"
    s[f"B{r}"] = f"=StartBank+SUM({_rng('Profit')})"
    s[f"B{r}"].number_format = "£#,##0.00"
    wb.defined_names["CurrentBank"] = DefinedName("CurrentBank", attr_text=f"Settings!$B${r}")
    s[f"A{r + 1}"] = "Current unit (£)"
    s[f"B{r + 1}"] = "=ROUND(CurrentBank*UnitPct,2)"
    s[f"B{r + 1}"].number_format = "£#,##0.00"
    wb.defined_names["Unit"] = DefinedName("Unit", attr_text=f"Settings!$B${r + 1}")
    s[f"A{r + 3}"] = "Yellow cells are yours to change. Everything else is worked out for you."
    s[f"A{r + 3}"].font = Font(italic=True, color="666666")

    # ------------------------------------------------------------ Bets
    b = wb.create_sheet("Bets")
    for i, (name, width) in enumerate(COLUMNS, start=1):
        c = b.cell(row=1, column=i, value=name)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
        c.alignment = Alignment(wrap_text=True, vertical="center")
        b.column_dimensions[get_column_letter(i)].width = width
    b.row_dimensions[1].height = 32
    b.freeze_panes = "C2"
    C = COL
    for row in range(2, MAX_ROWS + 2):
        g, h, d, j, l, m = (f"{C['Odds']}{row}", f"{C['Model chance']}{row}", f"{C['Type']}{row}",
                            f"{C['Edge']}{row}", f"{C['Stake']}{row}", f"{C['Result']}{row}")
        a = f"{C['Date']}{row}"
        b[f"{C['Fair odds']}{row}"] = f'=IF(OR({h}="",{h}=0),"",1/{h})'
        b[f"{C['Edge']}{row}"] = f'=IF(OR({g}="",{h}=""),"",{h}*{g}-1)'
        b[f"{C['Suggested stake']}{row}"] = (
            f'=IF(OR({g}="",{h}="",{g}<=1),"",'
            f'IF({j}<IF({d}="Builder",MinEdgeBuilder,IF({d}="Player",MinEdgePlayer,MinEdgeSingle)),0,'
            f'ROUND(MIN(IF({d}="Single",KellySingle,KellyBuilder)'
            f'*(StartBank+SUM({C["Profit"]}$1:{C["Profit"]}{row - 1}))*{j}/({g}-1),'
            f'IF({d}="Builder",CapBuilder,CapSingle)*Unit),2)))')
        b[f"{C['Return']}{row}"] = (
            f'=IF({l}="","",IF({m}="Won",{l}*{g},IF({m}="Void",{l},'
            f'IF({m}="Cash out",N({C["Cash-out return"]}{row}),IF({m}="Lost",0,"")))))')
        b[f"{C['Profit']}{row}"] = (
            f'=IF(OR({l}="",{m}="",{m}="Pending"),"",{C["Return"]}{row}-{l})')
        q = f"{C['Closing odds']}{row}"
        b[f"{C['CLV']}{row}"] = f'=IF(OR({g}="",{q}=""),"",{g}/{q}-1)'
        b[f"{C['Expected profit']}{row}"] = f'=IF(OR({l}="",{j}=""),"",{l}*{j})'
        b[f"{C['Bank after']}{row}"] = (
            f'=IF({C["Profit"]}{row}="","",StartBank+SUM({C["Profit"]}$2:{C["Profit"]}{row}))')
        b[a].number_format = "dd/mm/yyyy"
        for col in ("Model chance", "Edge", "CLV"):
            b[f"{C[col]}{row}"].number_format = "0.0%"
        for col in ("Fair odds", "Odds", "Closing odds"):
            b[f"{C[col]}{row}"].number_format = "0.00"
        for col in ("Suggested stake", "Stake", "Cash-out return", "Return", "Profit",
                    "Expected profit", "Bank after"):
            b[f"{C[col]}{row}"].number_format = "£#,##0.00"
    for col in ("Date", "Match", "Bookmaker", "Type", "Market", "Selection / legs", "Odds",
                "Model chance", "Stake", "Result", "Cash-out return", "Closing odds", "Notes"):
        for row in range(2, 60):
            b[f"{C[col]}{row}"].fill = INPUT_FILL

    def dv(col, options):
        v = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        v.add(f"{C[col]}2:{C[col]}{MAX_ROWS + 1}")
        b.add_data_validation(v)

    dv("Type", BET_TYPES)
    dv("Market", MARKETS)
    dv("Result", RESULTS)
    green = PatternFill("solid", fgColor="D9F2D9")
    red = PatternFill("solid", fgColor="F8D7D7")
    rng = f"{C['Profit']}2:{C['Profit']}{MAX_ROWS + 1}"
    b.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0"], fill=green))
    b.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0"], fill=red))

    # ------------------------------------------------------------ Summary
    sm = wb.create_sheet("Summary", 0)
    sm.column_dimensions["A"].width = 34
    for col in "BCDEFGH":
        sm.column_dimensions[col].width = 13
    sm["A1"] = "Betting record"
    sm["A1"].font = Font(bold=True, size=14)
    settled = f'{_rng("Result")},"<>Pending",{_rng("Result")},"<>"'
    stake_settled = f'SUMIFS({_rng("Stake")},{settled})'
    rows = [
        ("Bets placed", f'=COUNTA({_rng("Stake")})', "0"),
        ("Settled", f'=COUNTIFS({settled},{_rng("Stake")},"<>")', "0"),
        ("Pending", f'=COUNTIFS({_rng("Result")},"Pending")', "0"),
        ("Total staked (settled)", f"={stake_settled}", "£#,##0.00"),
        ("Total returned", f'=SUM({_rng("Return")})', "£#,##0.00"),
        ("Profit / loss", f'=SUM({_rng("Profit")})', "£#,##0.00"),
        ("ROI (profit ÷ staked)", f'=IF({stake_settled}=0,"",SUM({_rng("Profit")})/{stake_settled})', "0.0%"),
        ("Strike rate (won ÷ won+lost)",
         f'=IFERROR(COUNTIFS({_rng("Result")},"Won")/(COUNTIFS({_rng("Result")},"Won")+COUNTIFS({_rng("Result")},"Lost")),"")', "0.0%"),
        ("Average odds (settled)", f'=IFERROR(SUMPRODUCT(({_rng("Result")}="Won")+({_rng("Result")}="Lost"),{_rng("Odds")})/(COUNTIFS({_rng("Result")},"Won")+COUNTIFS({_rng("Result")},"Lost")),"")', "0.00"),
        ("Expected profit (model)", f'=SUMIFS({_rng("Expected profit")},{settled})', "£#,##0.00"),
        ("Actual minus expected", "=B8-B12", "£#,##0.00"),
        ("Average closing-line value", f'=IFERROR(AVERAGE({_rng("CLV")}),"")', "0.0%"),
        ("Bets that beat the closing odds", f'=IFERROR(COUNTIF({_rng("CLV")},">0")/COUNT({_rng("CLV")}),"")', "0.0%"),
        ("Current bank", "=CurrentBank", "£#,##0.00"),
        ("Bank change since start", '=IF(StartBank=0,"",CurrentBank/StartBank-1)', "0.0%"),
    ]
    for i, (label, formula, fmt) in enumerate(rows, start=3):
        sm[f"A{i}"] = label
        sm[f"B{i}"] = formula
        sm[f"B{i}"].number_format = fmt
        sm[f"B{i}"].font = Font(bold=label in ("Profit / loss", "ROI (profit ÷ staked)"))
    r0 = 3 + len(rows) + 1
    sm[f"A{r0}"] = ("ROI is your real return per £ staked. Closing-line value (CLV) shows whether "
                    "you are getting better odds than the final market price; over hundreds of bets "
                    "it is the best early sign of a genuine edge.")
    sm[f"A{r0}"].alignment = Alignment(wrap_text=True)
    sm.merge_cells(f"A{r0}:H{r0}")
    sm.row_dimensions[r0].height = 45

    def breakdown(r, title, key_col, keys):
        sm[f"A{r}"] = title
        sm[f"A{r}"].font = Font(bold=True)
        heads = ["Bets", "Staked", "Profit", "ROI", "Won", "Strike rate"]
        for j, h in enumerate(heads):
            c = sm.cell(row=r, column=2 + j, value=h)
            c.font = Font(bold=True)
        for k, key in enumerate(keys, start=1):
            rr = r + k
            sm[f"A{rr}"] = key
            crit = f'{_rng(key_col)},$A{rr},{_rng("Result")},"<>Pending",{_rng("Result")},"<>"'
            sm[f"B{rr}"] = f"=COUNTIFS({crit})"
            sm[f"C{rr}"] = f'=SUMIFS({_rng("Stake")},{crit})'
            sm[f"D{rr}"] = f'=SUMIFS({_rng("Profit")},{crit})'
            sm[f"E{rr}"] = f'=IF(C{rr}=0,"",D{rr}/C{rr})'
            sm[f"F{rr}"] = f'=COUNTIFS({_rng(key_col)},$A{rr},{_rng("Result")},"Won")'
            sm[f"G{rr}"] = (f'=IFERROR(F{rr}/(F{rr}+COUNTIFS({_rng(key_col)},$A{rr},'
                            f'{_rng("Result")},"Lost")),"")')
            sm[f"C{rr}"].number_format = sm[f"D{rr}"].number_format = "£#,##0.00"
            sm[f"E{rr}"].number_format = sm[f"G{rr}"].number_format = "0.0%"
        return r + len(keys) + 2

    r = breakdown(r0 + 2, "By bet type", "Type", BET_TYPES)
    r = breakdown(r, "By market", "Market", MARKETS)

    # Monthly
    sm[f"A{r}"] = "By month"
    sm[f"A{r}"].font = Font(bold=True)
    for j, h in enumerate(["Bets", "Staked", "Profit", "ROI", "Bank at month end"]):
        sm.cell(row=r, column=2 + j, value=h).font = Font(bold=True)
    first = dt.date(start.year, start.month, 1)
    for k in range(24):
        rr = r + 1 + k
        y, mth = divmod(first.month - 1 + k, 12)
        d0 = dt.date(first.year + y, mth + 1, 1)
        sm[f"A{rr}"] = d0
        sm[f"A{rr}"].number_format = "mmm yyyy"
        crit = (f'{_rng("Date")},">="&$A{rr},{_rng("Date")},"<"&EDATE($A{rr},1),'
                f'{_rng("Result")},"<>Pending",{_rng("Result")},"<>"')
        sm[f"B{rr}"] = f"=COUNTIFS({crit})"
        sm[f"C{rr}"] = f'=SUMIFS({_rng("Stake")},{crit})'
        sm[f"D{rr}"] = f'=SUMIFS({_rng("Profit")},{crit})'
        sm[f"E{rr}"] = f'=IF(C{rr}=0,"",D{rr}/C{rr})'
        sm[f"F{rr}"] = (f'=StartBank+SUMIFS({_rng("Profit")},{_rng("Date")},"<"&EDATE($A{rr},1))')
        for col in "CDF":
            sm[f"{col}{rr}"].number_format = "£#,##0.00"
        sm[f"E{rr}"].number_format = "0.0%"
    month_rows = (r + 1, r + 24)

    chart = LineChart()
    chart.title = "Bank at month end"
    chart.height, chart.width = 7, 16
    chart.y_axis.title = "£"
    data = Reference(sm, min_col=6, min_row=r, max_row=month_rows[1])
    cats = Reference(sm, min_col=1, min_row=month_rows[0], max_row=month_rows[1])
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    sm.add_chart(chart, "J3")

    # ------------------------------------------------------------ Guide
    g = wb.create_sheet("How to use")
    g.column_dimensions["A"].width = 110
    guide = [
        "How to use this tracker",
        "",
        "1. Settings: enter your starting bank (money set aside for betting only). Leave the rest unless you want to change the plan.",
        "2. Bets: add one row per bet before the match. Fill in the yellow cells: date, match, bookmaker, type (Single, Player or Builder),",
        "   market, what you backed, the odds you took, the model's chance (from the match report) and your stake.",
        "   'Suggested stake' shows what the staking plan recommends. 0 means the edge is too small: skip the bet.",
        "3. After the match set Result to Won, Lost, Void or Cash out (and enter the cash-out amount).",
        "4. Optional but valuable: enter the closing odds (the price just before kick-off) so the sheet can work out CLV.",
        "5. Summary updates itself: profit, ROI, strike rate, model-expected profit versus actual, and results by month,",
        "   bet type and market.",
        "",
        "Reading the numbers",
        "- ROI = profit ÷ total staked. Bookmaker margins mean the average bettor's ROI is negative; anything above 0% over",
        "  several hundred bets is good. Over 50 bets, luck dominates: do not judge the model on a handful of results.",
        "- Actual minus expected: positive means you have run better than the model expected, negative worse.",
        "- CLV above 0% on most bets means you are consistently getting better prices than the final market. That is the",
        "  clearest early sign the edge is real.",
        "",
        "You can also just tell Claude in the project what you backed and the result, and it will fill the sheet in for you.",
    ]
    for i, line in enumerate(guide, start=1):
        g[f"A{i}"] = line
        if i == 1 or line == "Reading the numbers":
            g[f"A{i}"].font = Font(bold=True, size=12)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


# --------------------------------------------------------------- editing


def _first_empty_row(ws) -> int:
    col = COL["Date"]
    for row in range(2, MAX_ROWS + 2):
        if ws[f"{col}{row}"].value in (None, ""):
            return row
    raise RuntimeError("tracker is full")


def add_bet(path: str | Path, *, date, match: str, selection: str, odds: float, stake: float,
            model_chance: float | None = None, bet_type: str = "Single", market: str = "Other",
            bookmaker: str = "", result: str = "Pending", closing_odds: float | None = None,
            notes: str = "") -> int:
    """Append a bet. Returns the sheet row number."""
    wb = load_workbook(path)
    ws = wb["Bets"]
    row = _first_empty_row(ws)
    vals = {
        "Date": pd.Timestamp(date).to_pydatetime(), "Match": match, "Bookmaker": bookmaker,
        "Type": bet_type, "Market": market, "Selection / legs": selection, "Odds": float(odds),
        "Model chance": model_chance, "Stake": float(stake), "Result": result,
        "Closing odds": closing_odds, "Notes": notes,
    }
    for k, v in vals.items():
        if v is not None:
            ws[f"{COL[k]}{row}"] = v
    wb.save(path)
    return row


def settle(path: str | Path, row: int, result: str, closing_odds: float | None = None,
           cash_out: float | None = None) -> None:
    if result not in RESULTS:
        raise ValueError(f"result must be one of {RESULTS}")
    wb = load_workbook(path)
    ws = wb["Bets"]
    ws[f"{COL['Result']}{row}"] = result
    if closing_odds is not None:
        ws[f"{COL['Closing odds']}{row}"] = float(closing_odds)
    if cash_out is not None:
        ws[f"{COL['Cash-out return']}{row}"] = float(cash_out)
    wb.save(path)


def read_bets(path: str | Path) -> pd.DataFrame:
    """The raw bet rows (inputs only), with profit computed in Python."""
    wb = load_workbook(path, data_only=False)
    ws = wb["Bets"]
    recs = []
    for row in range(2, MAX_ROWS + 2):
        if ws[f"{COL['Date']}{row}"].value in (None, ""):
            continue
        r = {k: ws[f"{COL[k]}{row}"].value for k in (
            "Date", "Match", "Bookmaker", "Type", "Market", "Selection / legs", "Odds",
            "Model chance", "Stake", "Result", "Cash-out return", "Closing odds", "Notes")}
        r["row"] = row
        recs.append(r)
    df = pd.DataFrame(recs)
    if df.empty:
        return df
    ret = []
    for _, r in df.iterrows():
        res, st, o = r["Result"], r["Stake"] or 0, r["Odds"] or 0
        ret.append(st * o if res == "Won" else st if res == "Void" else
                   (r["Cash-out return"] or 0) if res == "Cash out" else 0 if res == "Lost" else None)
    df["Return"] = ret
    df["Profit"] = [None if rt is None else rt - (st or 0) for rt, st in zip(df["Return"], df["Stake"])]
    return df


def summary(path: str | Path) -> dict:
    wb = load_workbook(path, data_only=False)
    start_bank = wb["Settings"]["B3"].value
    df = read_bets(path)
    if df.empty:
        return {"bets": 0, "bank": start_bank}
    s = df[df["Profit"].notna()]
    staked = float(s["Stake"].sum())
    profit = float(s["Profit"].sum())
    won = int((s["Result"] == "Won").sum())
    lost = int((s["Result"] == "Lost").sum())
    clv = (df["Odds"] / df["Closing odds"] - 1).dropna() if "Closing odds" in df else pd.Series()
    exp = (s["Stake"] * (s["Model chance"].astype(float) * s["Odds"] - 1)).dropna()
    return {
        "bets": len(df), "settled": len(s), "pending": int((df["Result"] == "Pending").sum()),
        "staked": staked, "profit": profit, "roi": profit / staked if staked else None,
        "strike_rate": won / (won + lost) if won + lost else None,
        "expected_profit": float(exp.sum()) if len(exp) else None,
        "avg_clv": float(clv.mean()) if len(clv) else None,
        "bank": (start_bank or 0) + profit,
    }
